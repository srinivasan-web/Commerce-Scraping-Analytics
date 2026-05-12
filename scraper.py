import argparse
import os
import time
import random
import logging
from datetime import datetime
from pathlib import Path
import re
from urllib.parse import urljoin, urlparse

import pandas as pd
from fake_useragent import UserAgent
from playwright.sync_api import sync_playwright, TimeoutError


# =========================================================
# LOGGING CONFIGURATION
# =========================================================

logging.basicConfig(
    filename="logs/scraper.log",
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)


# =========================================================
# GLOBAL CONFIG
# =========================================================

DEFAULT_URL = "https://www.amazon.in/gp/bestsellers/sports/3403930031/ref=zg_bs_nav_sports_2_206211218031"
BASE_URL = os.getenv("TARGET_URL", DEFAULT_URL)
MARKETPLACE_BASE_URL = "https://www.amazon.in"

HEADLESS_MODE = False

MAX_RETRIES = 3
MAX_PAGES = int(os.getenv("MAX_PAGES", "1"))
MAX_PRODUCTS = int(os.getenv("MAX_PRODUCTS", "30"))

OUTPUT_CSV = "output/amazon_bestsellers.csv"
OUTPUT_XLSX = "output/amazon_bestsellers.xlsx"

REMOVED_EXPORT_COLUMNS = {
    "Limited Time Deal",
    "Sponsored Status",
}


def timestamped_output_path(path):
    """
    Return a same-folder fallback path when the normal export file is locked.
    """
    output_path = Path(path)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return str(output_path.with_name(f"{output_path.stem}_{timestamp}{output_path.suffix}"))


def format_simple_excel(path):
    """
    Apply formatting to the single-sheet scraper workbook.
    """
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, PatternFill, Font

    workbook = load_workbook(path)
    worksheet = workbook.active

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for column in worksheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        header = clean_text(column[0].value)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 50)
        worksheet.column_dimensions[column_letter].width = adjusted_width

        if header in {"Estimated Monthly Revenue", "Product Total Revenue"}:
            for cell in column[1:]:
                cell.number_format = '"INR" #,##0.00'
        elif header in {"Overall Bought Count", "Reviews Count"}:
            for cell in column[1:]:
                cell.number_format = '#,##0'
        elif header == "Rank Number":
            for cell in column[1:]:
                cell.number_format = '0'
        elif header in {"Discount Percentage", "Product Rating Value"}:
            for cell in column[1:]:
                cell.number_format = '0.00'

    workbook.save(path)


def write_simple_excel(df, output_path):
    """
    Write the single-sheet Excel workbook and apply formatting.
    """
    df.to_excel(
        output_path,
        index=False,
        engine="openpyxl"
    )
    format_simple_excel(output_path)


# =========================================================
# DATA CLEANING AND ANALYTICS HELPERS
# =========================================================

def clean_text(value):
    """
    Normalize text scraped from Amazon cards.
    """
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def infer_marketplace_base_url(url):
    """
    Infer marketplace origin from the provided Amazon URL.
    """
    parsed = urlparse(clean_text(url))
    if parsed.scheme and parsed.netloc:
        return f"{parsed.scheme}://{parsed.netloc}"
    return MARKETPLACE_BASE_URL


def extract_asin(value):
    """
    Extract ASIN from common Amazon URLs or text.
    """
    match = re.search(r"(?:/dp/|/gp/product/|asin=)([A-Z0-9]{10})", clean_text(value))
    return match.group(1) if match else ""


def is_product_url(url):
    """
    Detect whether the provided URL points directly to an Amazon product page.
    """
    return bool(extract_asin(url))


def parse_price(value):
    """
    Convert price text like '₹1,299' into 1299.0.
    """
    text = clean_text(value)
    match = re.search(r"[\d,]+(?:\.\d+)?", text)
    if not match:
        return 0.0
    return float(match.group(0).replace(",", ""))


def parse_int(value):
    """
    Convert text containing an integer into int.
    """
    text = clean_text(value)
    match = re.search(r"\d[\d,]*", text)
    if not match:
        return 0
    return int(match.group(0).replace(",", ""))


def parse_rank(value, fallback=0):
    """
    Convert rank text like '#12' into 12.
    """
    rank = parse_int(value)
    return rank or fallback


def parse_bought_count(value):
    """
    Convert Amazon text like '500+ bought in past month' or '1K+ bought' to int.
    """
    text = clean_text(value).lower().replace(",", "")
    match = re.search(r"(\d+(?:\.\d+)?)\s*([km]?)\+?\s*bought", text)
    if not match:
        return 0

    number = float(match.group(1))
    suffix = match.group(2)
    if suffix == "k":
        number *= 1000
    elif suffix == "m":
        number *= 1_000_000
    return int(number)


def extract_card_metrics(full_text):
    """
    Extract analysis fields available directly in a bestseller product card.
    """
    text = clean_text(full_text)
    lines = [clean_text(line) for line in full_text.splitlines() if clean_text(line)]
    lowered = text.lower()

    price_matches = re.findall(r"(?:₹|â‚¹|INR|Rs\.?)\s*[\d,]+(?:\.\d+)?", text, flags=re.IGNORECASE)
    current_price = price_matches[0] if price_matches else ""
    original_price = price_matches[1] if len(price_matches) > 1 else ""

    discount_text = ""
    discount_match = re.search(r"(\d+(?:\.\d+)?)\s*%\s*(?:off|discount)?", text, flags=re.IGNORECASE)
    if discount_match:
        discount_text = f"{discount_match.group(1)}%"

    bought_text = ""
    bought_match = re.search(r"(\d[\d,.]*\s*[kKmM]?\+?\s*bought\s+in\s+past\s+month)", text, flags=re.IGNORECASE)
    if bought_match:
        bought_text = clean_text(bought_match.group(1))

    rating = ""
    rating_match = re.search(r"([0-5](?:\.\d+)?)\s*out of 5", text, flags=re.IGNORECASE)
    if rating_match:
        rating = rating_match.group(1)

    reviews = ""
    reviews_match = re.search(r"out of 5 stars\s+([\d,]+)", text, flags=re.IGNORECASE)
    if reviews_match:
        reviews = reviews_match.group(1)

    price_value = parse_price(current_price)
    original_price_value = parse_price(original_price)
    discount_pct = float(discount_match.group(1)) if discount_match else 0.0
    if not discount_pct and original_price_value > price_value > 0:
        discount_pct = round((original_price_value - price_value) / original_price_value * 100, 2)

    overall_bought_count = parse_bought_count(bought_text)
    estimated_revenue = round(price_value * overall_bought_count, 2)

    return {
        "current_price_text": current_price,
        "current_price_value": price_value,
        "original_price_text": original_price,
        "original_price_value": original_price_value,
        "discount_text": discount_text,
        "discount_pct": discount_pct,
        "rating": rating,
        "rating_value": float(rating) if rating else 0.0,
        "reviews": reviews,
        "reviews_count": parse_int(reviews),
        "bought_text": bought_text,
        "overall_bought_count": overall_bought_count,
        "estimated_monthly_revenue": estimated_revenue,
        "product_total_revenue": estimated_revenue,
        "is_prime": "prime" in lowered,
        "free_delivery": "free delivery" in lowered or "free" in lowered,
        "limited_time_deal": "limited time deal" in lowered or "deal" in lowered,
        "coupon_available": "coupon" in lowered,
        "sponsored_status": "sponsored" in lowered,
        "bestseller_badge": "best seller" in lowered or "#1 best seller" in lowered,
        "availability": "Out of Stock" if "out of stock" in lowered else "Available",
        "card_text": " | ".join(lines[:25]),
    }


# =========================================================
# RANDOM USER AGENT
# =========================================================

def get_random_user_agent():
    """
    Generate random user-agent for avoiding bot detection
    """
    ua = UserAgent()
    return ua.random


# =========================================================
# RANDOM DELAY
# =========================================================

def random_delay(min_sec=1, max_sec=3):
    """
    Random sleep to simulate human behavior
    """
    delay = random.uniform(min_sec, max_sec)
    time.sleep(delay)


# =========================================================
# SCROLL PAGE
# =========================================================

def scroll_page(page, target_count=MAX_PRODUCTS):
    """
    Scroll page slowly to load dynamic content
    """
    print("Scrolling page to load products...")
    
    previous_height = 0
    scroll_count = 0
    max_scrolls = 6

    while scroll_count < max_scrolls:
        try:
            loaded_count = page.locator("div[data-asin]").count()
            if target_count and loaded_count >= target_count:
                break
        except:
            pass

        current_height = page.evaluate(
            "() => document.body.scrollHeight"
        )

        if current_height == previous_height:
            break

        previous_height = current_height
        page.mouse.wheel(0, 2000)
        random_delay(0.5, 1)
        scroll_count += 1
    
    print(f"Scrolled {scroll_count} times")


# =========================================================
# EXTRACT PRODUCT DATA FROM CONTAINER
# =========================================================

def extract_product_data(product_container, rank, category_name):
    """
    Extract product details from product container element
    """
    try:
        try:
            full_text = product_container.inner_text()
        except:
            full_text = ""
        metrics = extract_card_metrics(full_text)

        # Get ASIN
        asin = None
        try:
            asin = product_container.get_attribute("data-asin")
        except:
            pass

        # Product Name - second link contains the product name
        product_name = None
        product_url = None
        try:
            # Get all links in the container
            links = product_container.locator("a")
            if links.count() >= 2:
                # The second link usually has the product name
                link = links.nth(1)
                product_name = link.inner_text().strip()
                href = link.get_attribute("href")
                if href:
                    if not href.startswith("http"):
                        product_url = "https://www.amazon.in" + href
                    else:
                        product_url = href
        except:
            pass

        # If name still not found, try inner text
        if not product_name:
            try:
                full_text = product_container.inner_text()
                # Extract lines and find product name
                lines = [l.strip() for l in full_text.split('\n') if l.strip() and not l.startswith('#')]
                if len(lines) > 0:
                    product_name = lines[0]
            except:
                pass

        # Brand (first word of product name)
        brand = None
        if product_name:
            words = product_name.split()
            if words:
                brand = words[0]

        # Try to get price - it's in the text as ₹XXX.XX format
        price = None
        try:
            full_text = product_container.inner_text()
            import re
            price_match = re.search(r'₹\s*[\d,]+\.?\d*', full_text)
            if price_match:
                price = price_match.group(0)
        except:
            pass

        # Try to get rating - look for "X.X out of 5 stars" pattern
        rating = None
        try:
            full_text = product_container.inner_text()
            import re
            rating_match = re.search(r'(\d\.?\d*)\s*out of 5', full_text)
            if rating_match:
                rating = rating_match.group(1)
        except:
            pass

        # Try to get number of reviews
        reviews = None
        try:
            # The reviews count follows the rating
            full_text = product_container.inner_text()
            import re
            # Look for pattern: "4.2 out of 5 stars" followed by a number
            reviews_match = re.search(r'out of 5 stars\s+([\d,]+)', full_text)
            if reviews_match:
                reviews = reviews_match.group(1)
        except:
            pass

        price = metrics["current_price_text"] or price
        rating = metrics["rating"] or rating
        reviews = metrics["reviews"] or reviews

        # Get image URL
        image_url = None
        try:
            img = product_container.locator("img").first
            if img.count() > 0:
                image_url = img.get_attribute("src")
        except:
            pass

        # Timestamp
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        data = {
            "Product Rank": rank,
            "Rank Number": parse_rank(rank),
            "Product Name": product_name,
            "Brand Name": brand,
            "Product Price": price,
            "Product Price Value": metrics["current_price_value"],
            "Original Price": metrics["original_price_text"],
            "Original Price Value": metrics["original_price_value"],
            "Discount Percentage": metrics["discount_pct"],
            "Discount Text": metrics["discount_text"],
            "Product Rating": rating,
            "Product Rating Value": metrics["rating_value"],
            "Number of Reviews": reviews,
            "Reviews Count": metrics["reviews_count"],
            "Overall Bought Count": metrics["overall_bought_count"],
            "Bought Count Text": metrics["bought_text"],
            "Estimated Monthly Revenue": metrics["estimated_monthly_revenue"],
            "Product Total Revenue": metrics["product_total_revenue"],
            "Product URL": product_url,
            "Product Image URL": image_url,
            "Availability": metrics["availability"],
            "Category": category_name,
            "ASIN": asin,
            "Prime Available": metrics["is_prime"],
            "Free Delivery": metrics["free_delivery"],
            "Limited Time Deal": metrics["limited_time_deal"],
            "Coupon Available": metrics["coupon_available"],
            "Sponsored Status": metrics["sponsored_status"],
            "Bestseller Badge": metrics["bestseller_badge"],
            "Raw Card Text": metrics["card_text"],
            "Timestamp": timestamp
        }

        return data

    except Exception as e:
        logging.error(f"Error extracting product: {e}")
        return None


def build_product_row(raw_product, category_name, fallback_rank):
    """
    Build an analysis-ready row from raw card data extracted in one browser call.
    """
    full_text = raw_product.get("text", "")
    metrics = extract_card_metrics(full_text)

    product_name = clean_text(raw_product.get("name"))
    brand = product_name.split()[0] if product_name else ""
    rank = clean_text(raw_product.get("rank")) or f"#{fallback_rank}"
    rank_number = parse_rank(rank, fallback_rank)
    product_url = clean_text(raw_product.get("url"))
    image_url = clean_text(raw_product.get("image"))
    asin = clean_text(raw_product.get("asin"))

    return {
        "Product Rank": rank,
        "Rank Number": rank_number,
        "Product Name": product_name,
        "Brand Name": brand,
        "Product Price": metrics["current_price_text"],
        "Product Price Value": metrics["current_price_value"],
        "Original Price": metrics["original_price_text"],
        "Original Price Value": metrics["original_price_value"],
        "Discount Percentage": metrics["discount_pct"],
        "Discount Text": metrics["discount_text"],
        "Product Rating": metrics["rating"],
        "Product Rating Value": metrics["rating_value"],
        "Number of Reviews": metrics["reviews"],
        "Reviews Count": metrics["reviews_count"],
        "Overall Bought Count": metrics["overall_bought_count"],
        "Bought Count Text": metrics["bought_text"],
        "Estimated Monthly Revenue": metrics["estimated_monthly_revenue"],
        "Product Total Revenue": metrics["product_total_revenue"],
        "Product URL": product_url,
        "Product Image URL": image_url,
        "Availability": metrics["availability"],
        "Category": category_name,
        "ASIN": asin,
        "Prime Available": metrics["is_prime"],
        "Free Delivery": metrics["free_delivery"],
        "Limited Time Deal": metrics["limited_time_deal"],
        "Coupon Available": metrics["coupon_available"],
        "Sponsored Status": metrics["sponsored_status"],
        "Bestseller Badge": metrics["bestseller_badge"],
        "Raw Card Text": metrics["card_text"],
        "Timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def extract_products_fast(page, category_name, max_products=MAX_PRODUCTS):
    """
    Extract all visible product cards in one browser-side pass for better speed.
    """
    raw_products = page.evaluate(
        """
        ({ marketplaceBaseUrl, maxProducts }) => {
            const clean = (value) => (value || "").replace(/\\s+/g, " ").trim();
            const absolute = (href) => {
                try {
                    return href ? new URL(href, marketplaceBaseUrl).href : "";
                } catch {
                    return "";
                }
            };
            const cards = Array.from(document.querySelectorAll("div[data-asin]"))
                .filter((card) => clean(card.getAttribute("data-asin")) || card.querySelector("a[href*='/dp/']"));

            return cards.slice(0, maxProducts || cards.length).map((card, index) => {
                const text = card.innerText || "";
                const link = card.querySelector("a[href*='/dp/'], a[href*='/gp/product/']");
                const img = card.querySelector("img");
                const rankText = clean(card.querySelector(".zg-bdg-text, [class*='zg-bdg-text']")?.innerText);
                const name =
                    clean(card.querySelector("div[class*='p13n-sc-css-line-clamp']")?.innerText) ||
                    clean(card.querySelector("a.a-link-normal div")?.innerText) ||
                    clean(img?.getAttribute("alt")) ||
                    clean(link?.innerText);

                return {
                    index: index + 1,
                    rank: rankText || `#${index + 1}`,
                    asin: clean(card.getAttribute("data-asin")),
                    name,
                    url: absolute(link?.getAttribute("href")),
                    image: clean(img?.getAttribute("src") || img?.getAttribute("data-src")),
                    text,
                };
            }).filter((item) => item.name && item.url);
        }
        """,
        {"marketplaceBaseUrl": MARKETPLACE_BASE_URL, "maxProducts": max_products},
    )

    products = []
    seen = set()
    for index, raw_product in enumerate(raw_products, start=1):
        product = build_product_row(raw_product, category_name, index)
        key = product["ASIN"] or product["Product URL"] or product["Product Name"]
        if key in seen:
            continue
        seen.add(key)
        products.append(product)
        if max_products and len(products) >= max_products:
            break
    return products


def extract_product_detail_live(page, category_name, source_url):
    """
    Scrape a direct Amazon product page when there are no bestseller cards.
    """
    raw_product = page.evaluate(
        """
        ({ marketplaceBaseUrl }) => {
            const clean = (value) => (value || "").replace(/\\s+/g, " ").trim();
            const text = (selector) => clean(document.querySelector(selector)?.innerText || "");
            const attr = (selector, name) => clean(document.querySelector(selector)?.getAttribute(name) || "");
            const absolute = (href) => {
                try {
                    return href ? new URL(href, marketplaceBaseUrl).href : window.location.href;
                } catch {
                    return window.location.href;
                }
            };

            const title = text("#productTitle") || text("span#title") || text("h1");
            const asin = attr("[name='ASIN']", "value");
            const image = attr("#landingImage", "src") || attr("#imgTagWrapperId img", "src");

            return {
                index: 1,
                rank: "#1",
                asin,
                name: title,
                url: absolute(window.location.href),
                image,
                text: document.body?.innerText || "",
            };
        }
        """,
        {"marketplaceBaseUrl": MARKETPLACE_BASE_URL},
    )

    if not clean_text(raw_product.get("name")):
        return []
    product = build_product_row(raw_product, category_name, 1)
    product["Product URL"] = source_url or product["Product URL"]
    return [product]


# =========================================================
# SCRAPE SINGLE PAGE
# =========================================================

def scrape_page(page, max_products=MAX_PRODUCTS):
    """
    Scrape products from current page
    """
    all_products = []

    try:
        # Wait for product containers to load
        try:
            page.wait_for_selector("div[data-asin]", timeout=15000)
        except:
            print("Warning: Timeout waiting for product containers, continuing anyway...")
        
        # Get category
        category_name = "Sports & Fitness"
        try:
            category_name = page.locator("h1").first.inner_text().strip()
        except:
            pass

        # Scroll to load all products
        scroll_page(page, target_count=max_products)

        # Get all product containers using data-asin attribute
        product_containers = page.locator("div[data-asin]")
        
        container_count = product_containers.count()
        print(f"Found {container_count} product containers on page")
        logging.info(f"Products Found: {container_count}")

        if container_count == 0:
            print("Warning: No containers found, trying alternative selectors...")
            # Try alternative approach
            all_divs = page.locator("div")
            print(f"Total divs: {all_divs.count()}")

        fast_products = extract_products_fast(page, category_name, max_products=max_products)
        if fast_products:
            print(f"Fast extraction captured {len(fast_products)} analysis-ready products")
            return fast_products

        if is_product_url(page.url):
            detail_products = extract_product_detail_live(page, category_name, page.url)
            if detail_products:
                print("Direct product page captured as live product detail")
                return detail_products

        for i in range(min(container_count, max_products)):
            try:
                container = product_containers.nth(i)
                rank = f"#{i + 1}"
                
                data = extract_product_data(container, rank, category_name)
                
                if data and data["Product Name"]:  # Only add if we got a name
                    all_products.append(data)
                
                random_delay(0.1, 0.3)

            except Exception as e:
                logging.error(f"Error processing product {i}: {e}")
                continue

    except TimeoutError:
        logging.error("Timeout while loading page")
        print("Timeout while loading page")

    print(f"Successfully extracted {len(all_products)} products")
    return all_products


# =========================================================
# PAGINATION SUPPORT
# =========================================================

def go_to_next_page(page):
    """
    Navigate to next page if available
    """
    try:
        selectors = [
            "a[rel='next']",
            "li.a-last:not(.a-disabled) a",
            "a[aria-label='Next page']",
        ]

        for selector in selectors:
            next_buttons = page.locator(selector)
            if next_buttons.count() == 0:
                continue
            next_button = next_buttons.first
            href = next_button.get_attribute("href")
            if href:
                page.goto(urljoin(MARKETPLACE_BASE_URL, href), timeout=60000, wait_until="domcontentloaded")
            else:
                next_button.click()
            random_delay(2, 4)
            return True

        return False

    except Exception as e:
        logging.error(f"Pagination Error: {e}")
        return False


# =========================================================
# MAIN SCRAPER FUNCTION
# =========================================================

def run_scraper(target_url=BASE_URL, max_pages=MAX_PAGES, max_products=MAX_PRODUCTS, headless=HEADLESS_MODE):
    """
    Main scraper function with retry logic
    """
    all_data = []
    retry_count = 0
    target_url = clean_text(target_url) or BASE_URL

    while retry_count < MAX_RETRIES:
        try:
            print("Initializing Playwright...")
            with sync_playwright() as p:
                
                browser = p.chromium.launch(
                    headless=headless,
                    slow_mo=50
                )

                context = browser.new_context(
                    user_agent=get_random_user_agent(),
                    viewport={"width": 1400, "height": 900}
                )

                page = context.new_page()

                print(f"Opening live URL: {target_url}")
                logging.info(f"Opening live URL: {target_url}")

                page.goto(target_url, timeout=60000, wait_until="domcontentloaded")

                random_delay(2, 4)

                current_page = 1

                while True:
                    print(f"\n{'='*50}")
                    print(f"Scraping Page {current_page}")
                    print(f"{'='*50}")
                    logging.info(f"Scraping Page {current_page}")

                    remaining_products = max(max_products - len(all_data), 0) if max_products else MAX_PRODUCTS
                    if max_products and remaining_products <= 0:
                        print(f"Reached top {max_products} product limit")
                        break

                    page_data = scrape_page(page, max_products=remaining_products)
                    
                    print(f"Scraped {len(page_data)} products from page {current_page}")
                    all_data.extend(page_data)
                    if max_products and len(all_data) >= max_products:
                        all_data = all_data[:max_products]
                        print(f"Reached top {max_products} product limit")
                        break

                    if is_product_url(target_url):
                        print("Direct product URL scraped. Pagination skipped.")
                        break

                    if current_page >= max_pages:
                        print("Reached page limit")
                        break

                    # Check for next page
                    has_next = go_to_next_page(page)
                    
                    if not has_next:
                        print("No more pages available")
                        break

                    current_page += 1
                    
                browser.close()
                print("\nBrowser closed successfully")
                break

        except Exception as e:
            retry_count += 1
            print(f"\nError occurred (Retry {retry_count}/{MAX_RETRIES}): {str(e)}")
            logging.error(f"Retry {retry_count} Error: {e}")
            
            if retry_count < MAX_RETRIES:
                random_delay(5, 10)

    return all_data


# =========================================================
# SAVE DATA
# =========================================================

def save_data(data, max_products=MAX_PRODUCTS):
    """
    Save scraped data to CSV and Excel formats
    """
    if not data:
        print("❌ No data found to export")
        logging.error("No data found to export")
        return

    print(f"\n{'='*50}")
    print(f"Processing {len(data)} products...")
    print(f"{'='*50}")

    df = pd.DataFrame(data)

    numeric_columns = [
        "Product Price Value",
        "Original Price Value",
        "Discount Percentage",
        "Product Rating Value",
        "Reviews Count",
        "Overall Bought Count",
        "Estimated Monthly Revenue",
        "Product Total Revenue",
        "Rank Number",
    ]
    for column in numeric_columns:
        if column in df.columns:
            df[column] = pd.to_numeric(df[column], errors="coerce").fillna(0)

    # Remove duplicates
    initial_count = len(df)
    df.drop_duplicates(
        subset=["Product Name", "Product URL"],
        inplace=True
    )
    if "Rank Number" in df.columns:
        df.sort_values("Rank Number", ascending=True, inplace=True)
    if max_products:
        df = df.head(max_products)
    if "Product Total Revenue" not in df.columns and "Estimated Monthly Revenue" in df.columns:
        df["Product Total Revenue"] = df["Estimated Monthly Revenue"]
    df.drop(columns=[column for column in REMOVED_EXPORT_COLUMNS if column in df.columns], errors="ignore", inplace=True)

    preferred_columns = [
        "Product Rank",
        "Rank Number",
        "Product Name",
        "Brand Name",
        "Product Price",
        "Product Price Value",
        "Original Price",
        "Original Price Value",
        "Discount Percentage",
        "Product Rating",
        "Product Rating Value",
        "Number of Reviews",
        "Reviews Count",
        "Overall Bought Count",
        "Bought Count Text",
        "Estimated Monthly Revenue",
        "Product Total Revenue",
        "Availability",
        "Prime Available",
        "Free Delivery",
        "Coupon Available",
        "Bestseller Badge",
        "Product URL",
        "Product Image URL",
        "Category",
        "ASIN",
        "Raw Card Text",
        "Timestamp",
    ]
    ordered_columns = [column for column in preferred_columns if column in df.columns]
    remaining_columns = [column for column in df.columns if column not in ordered_columns]
    df = df[ordered_columns + remaining_columns]

    final_count = len(df)
    duplicates_removed = initial_count - final_count
    
    print(f"Total products: {initial_count}")
    print(f"After removing duplicates: {final_count}")
    print(f"Duplicates removed: {duplicates_removed}")

    global OUTPUT_CSV, OUTPUT_XLSX
    try:
        with open(OUTPUT_CSV, "a+b"):
            pass
    except PermissionError:
        OUTPUT_CSV = timestamped_output_path(OUTPUT_CSV)
        print(f"Default CSV is locked. Using fallback file: {OUTPUT_CSV}")

    try:
        with open(OUTPUT_XLSX, "a+b"):
            pass
    except PermissionError:
        OUTPUT_XLSX = timestamped_output_path(OUTPUT_XLSX)
        print(f"Default Excel workbook is locked. Using fallback file: {OUTPUT_XLSX}")

    # UTF-8 CSV Export
    try:
        df.to_csv(
            OUTPUT_CSV,
            index=False,
            encoding="utf-8-sig"
        )
        print(f"✓ CSV exported: {OUTPUT_CSV}")
        logging.info(f"CSV exported successfully: {OUTPUT_CSV}")
    except Exception as e:
        print(f"❌ Error exporting CSV: {e}")
        logging.error(f"Error exporting CSV: {e}")

    # Excel Export with formatting
    try:
        df.to_excel(
            OUTPUT_XLSX,
            index=False,
            engine="openpyxl"
        )
        
        # Add formatting to Excel
        try:
            from openpyxl import load_workbook
            from openpyxl.styles import Alignment, PatternFill, Font
            
            workbook = load_workbook(OUTPUT_XLSX)
            worksheet = workbook.active
            
            # Header formatting
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            # Auto-adjust column widths
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                header = clean_text(column[0].value)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width

                if header in {"Estimated Monthly Revenue", "Product Total Revenue"}:
                    for cell in column[1:]:
                        cell.number_format = '"INR" #,##0.00'
                elif header in {"Overall Bought Count", "Reviews Count"}:
                    for cell in column[1:]:
                        cell.number_format = '#,##0'
                elif header == "Rank Number":
                    for cell in column[1:]:
                        cell.number_format = '0'
                elif header in {"Discount Percentage", "Product Rating Value"}:
                    for cell in column[1:]:
                        cell.number_format = '0.00'
            
            workbook.save(OUTPUT_XLSX)
        except:
            pass  # If formatting fails, file still exists
        
        print(f"✓ Excel exported: {OUTPUT_XLSX}")
        logging.info(f"Excel exported successfully: {OUTPUT_XLSX}")
    except Exception as e:
        print(f"❌ Error exporting Excel: {e}")
        logging.error(f"Error exporting Excel: {e}")

    print(f"\n{'='*50}")
    print(f"✓ Data Export Complete!")
    print(f"{'='*50}")


# =========================================================
# MAIN
# =========================================================

def parse_args():
    """
    Parse runtime options for dynamic live scraping.
    """
    parser = argparse.ArgumentParser(description="Live scrape an Amazon bestseller/category URL or direct product URL.")
    parser.add_argument(
        "url",
        nargs="?",
        default=BASE_URL,
        help="Amazon URL to scrape. Can be a bestseller/category page or a direct /dp/ product page.",
    )
    parser.add_argument(
        "--max-pages",
        type=int,
        default=MAX_PAGES,
        help="Maximum category/search pages to scrape. Defaults to 1 for top bestseller scraping.",
    )
    parser.add_argument(
        "--max-products",
        type=int,
        default=MAX_PRODUCTS,
        help="Maximum products to scrape. Defaults to top 30.",
    )
    parser.add_argument(
        "--headless",
        action="store_true",
        help="Run Chromium in headless mode.",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    MARKETPLACE_BASE_URL = infer_marketplace_base_url(args.url)

    print("\n" + "="*50)
    print("Starting Amazon Live Scraper...")
    print("="*50 + "\n")
    print(f"Target URL: {args.url}")
    print(f"Marketplace: {MARKETPLACE_BASE_URL}")
    print(f"Max pages: {args.max_pages}")
    print(f"Max products: {args.max_products}")

    scraped_data = run_scraper(
        target_url=args.url,
        max_pages=args.max_pages,
        max_products=args.max_products,
        headless=args.headless,
    )

    save_data(scraped_data, max_products=args.max_products)

    print("\n✓ Scraping Completed.\n")
