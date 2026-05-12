"""
Simple and direct Amazon bestseller scraper with Excel export
"""
import time
import logging
from datetime import datetime
import re

import pandas as pd
from fake_useragent import UserAgent
from playwright.sync_api import sync_playwright

# Logging
logging.basicConfig(
    filename="logs/scraper.log",
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

BASE_URL = "https://www.amazon.in/gp/bestsellers/sports/3403930031/ref=zg_bs_nav_sports_2_206211218031"
OUTPUT_XLSX = "output/amazon_bestsellers.xlsx"
OUTPUT_CSV = "output/amazon_bestsellers.csv"

def scrape_amazon():
    """Scrape Amazon bestsellers"""
    print("\n" + "="*70)
    print("AMAZON BESTSELLER SCRAPER - EXCEL EXPORT")
    print("="*70 + "\n")
    
    ua = UserAgent()
    all_products = []
    
    try:
        with sync_playwright() as p:
            print("🔄 Launching browser...")
            browser = p.chromium.launch(headless=False)
            page = browser.new_page(user_agent=ua.random, viewport={"width": 1400, "height": 900})
            
            print("📍 Navigating to Amazon...")
            page.goto(BASE_URL, timeout=60000, wait_until="networkidle")
            time.sleep(3)
            
            print("📜 Scrolling page to load all products...")
            for i in range(10):
                page.mouse.wheel(0, 3000)
                time.sleep(0.3)
            
            time.sleep(2)
            
            print("🔍 Extracting product data using JavaScript...")
            
            # JavaScript to extract all products
            products = page.evaluate("""
            () => {
                const items = [];
                const containers = document.querySelectorAll('[data-asin]');
                
                containers.forEach((container, idx) => {
                    try {
                        const textContent = container.innerText || '';
                        const asin = container.getAttribute('data-asin');
                        
                        // Parse text to extract fields
                        const lines = textContent.split('\\n').map(l => l.trim()).filter(l => l);
                        
                        let rank = `#${idx + 1}`;
                        let name = '';
                        let price = '';
                        let rating = '';
                        let reviews = '';
                        
                        for (let i = 0; i < lines.length; i++) {
                            const line = lines[i];
                            
                            // Rank
                            if (line.match(/^#\\d+$/)) rank = line;
                            
                            // Price (has ₹)
                            if (line.includes('₹')) {
                                price = line.trim();
                            }
                            
                            // Rating (has "out of 5")
                            if (line.includes('out of 5')) {
                                const ratingMatch = line.match(/([0-9.]+)\\s+out of 5/);
                                if (ratingMatch) rating = ratingMatch[1];
                                
                                // Reviews are next
                                if (i + 1 < lines.length) {
                                    reviews = lines[i + 1].replace(/[^0-9,]/g, '');
                                }
                            }
                            
                            // Product name - first long line
                            if (!name && line.length > 15 && !line.includes('₹') && !line.includes('out of') && !line.match(/^#\\d+$/)) {
                                name = line;
                            }
                        }
                        
                        // Get URL
                        let url = '';
                        const link = container.querySelector('a[href*="/dp/"]');
                        if (link) {
                            let href = link.getAttribute('href');
                            url = href.startsWith('http') ? href : 'https://www.amazon.in' + href;
                        }
                        
                        // Get image
                        let image = '';
                        const img = container.querySelector('img');
                        if (img) image = img.getAttribute('src') || '';
                        
                        if (name && asin) {
                            items.push({
                                rank,
                                name,
                                brand: name.split(' ')[0],
                                price,
                                rating,
                                reviews,
                                url,
                                image,
                                asin
                            });
                        }
                    } catch (e) {
                        console.error('Error:', e);
                    }
                });
                
                return items;
            }
            """)
            
            print(f"✓ Found {len(products)} products\n")
            
            # Format for dataframe
            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            for p in products:
                all_products.append({
                    "Rank": p['rank'],
                    "Product Name": p['name'],
                    "Brand": p['brand'],
                    "Price": p['price'],
                    "Rating": p['rating'],
                    "Reviews": p['reviews'],
                    "Product URL": p['url'],
                    "Image URL": p['image'],
                    "ASIN": p['asin'],
                    "Category": "Sports & Fitness",
                    "Scraped On": timestamp
                })
            
            browser.close()
            print("Browser closed\n")
            
    except Exception as e:
        print(f"❌ Error: {e}")
        logging.error(f"Error: {e}")
        return []
    
    return all_products

def save_to_excel(products):
    """Save products to Excel"""
    if not products:
        print("❌ No products to save!")
        return
    
    print("="*70)
    print("SAVING TO EXCEL")
    print("="*70 + "\n")
    
    df = pd.DataFrame(products)
    
    print(f"Total products: {len(df)}")
    print(f"Columns: {', '.join(df.columns)}\n")
    
    # Save CSV
    try:
        df.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
        print(f"✓ CSV saved: {OUTPUT_CSV}")
    except Exception as e:
        print(f"❌ CSV Error: {e}")
    
    # Save Excel with formatting
    try:
        df.to_excel(OUTPUT_XLSX, index=False, engine="openpyxl")
        
        # Format Excel
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, PatternFill, Font, Border, Side
        
        wb = load_workbook(OUTPUT_XLSX)
        ws = wb.active
        
        # Header styling
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Format header
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Format data cells
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell.border = thin_border
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            col_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 3, 60)
            ws.column_dimensions[col_letter].width = adjusted_width
        
        # Set row height for header
        ws.row_dimensions[1].height = 25
        
        wb.save(OUTPUT_XLSX)
        print(f"✓ Excel saved: {OUTPUT_XLSX}")
        print(f"\n📊 Total Columns: {len(df.columns)}")
        print(f"📊 Total Rows: {len(df)}")
        
    except Exception as e:
        print(f"❌ Excel Error: {e}")
    
    print("\n" + "="*70)
    print("✓ EXPORT COMPLETE!")
    print("="*70 + "\n")
    
    # Print sample
    print("📋 SAMPLE DATA (First 3 products):\n")
    for i, row in df.head(3).iterrows():
        print(f"{i+1}. {row['Product Name'][:60]}")
        print(f"   Price: {row['Price']}")
        print(f"   Rating: {row['Rating']} ⭐ ({row['Reviews']} reviews)")
        print(f"   URL: {row['Product URL'][:70]}")
        print()

if __name__ == "__main__":
    products = scrape_amazon()
    save_to_excel(products)
