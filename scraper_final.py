import time
import logging
from datetime import datetime

import pandas as pd
from fake_useragent import UserAgent
from playwright.sync_api import sync_playwright

# =========================================================
# LOGGING
# =========================================================

logging.basicConfig(
    filename="logs/scraper.log",
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

# =========================================================
# CONFIG
# =========================================================

BASE_URL = "https://www.amazon.in/gp/bestsellers/sports/3403930031/ref=zg_bs_nav_sports_2_206211218031"
OUTPUT_CSV = "output/amazon_bestsellers.csv"
OUTPUT_XLSX = "output/amazon_bestsellers.xlsx"

# =========================================================
# SCRAPER
# =========================================================

def scrape_bestsellers():
    """Scrape Amazon bestsellers"""
    
    all_products = []
    
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        page = browser.new_page(user_agent=UserAgent().random)
        
        print("Opening URL...")
        page.goto(BASE_URL, wait_until="load")
        
        # Wait extra time for dynamic content
        time.sleep(5)
        
        # Scroll multiple times
        print("Scrolling...")
        for i in range(10):
            page.mouse.wheel(0, 2000)
            time.sleep(0.3)
        
        time.sleep(2)
        
        # Get products using a more reliable method
        print("Extracting products...")
        
        # Get all visible product links
        product_elements = page.query_selector_all('a[href*="/dp/"]')
        
        print(f"Found {len(product_elements)} product elements")
        
        rank = 1
        seen_names = set()
        
        for elem in product_elements:
            try:
                href = elem.get_attribute('href')
                text = elem.inner_text().strip()
                
                # Skip if no text or URL
                if not text or not href or len(text) < 5:
                    continue
                
                # Skip duplicate or non-product text
                if text.startswith('See') or text.startswith('Skip') or text.startswith('Your'):
                    continue
                
                # Skip ratings link (contains "out of 5")
                if "out of 5" in text:
                    continue
                
                # Skip if we've seen this product name
                if text in seen_names:
                    continue
                seen_names.add(text)
                
                # Create product URL
                if not href.startswith('http'):
                    product_url = "https://www.amazon.in" + href
                else:
                    product_url = href
                
                # Extract ASIN from URL
                import re
                asin_match = re.search(r'/dp/([A-Z0-9]+)', href)
                asin = asin_match.group(1) if asin_match else ""
                
                # Get parent container for more info
                try:
                    parent = elem
                    # Go up a few levels to get the product container
                    for _ in range(5):
                        parent_temp = page.evaluate('el => el.parentElement', parent)
                        if parent_temp:
                            parent = parent_temp
                        else:
                            break
                    parent_text = page.evaluate('el => el.innerText', parent) if parent else ""
                except:
                    parent_text = ""
                
                # Extract price if available
                price = ""
                if "₹" in parent_text:
                    price_match = re.search(r'₹\s*[\d,]+\.?\d*', parent_text)
                    if price_match:
                        price = price_match.group(0)
                
                # Extract rating if available
                rating = ""
                if "out of 5" in parent_text:
                    rating_match = re.search(r'(\d\.?\d*)\s*out of 5', parent_text)
                    if rating_match:
                        rating = rating_match.group(1)
                
                # Extract reviews if available
                reviews = ""
                if "out of 5" in parent_text:
                    reviews_match = re.search(r'out of 5 stars\s+([\d,]+)', parent_text)
                    if reviews_match:
                        reviews = reviews_match.group(1)
                
                # Brand (first word)
                brand = text.split()[0] if text else ""
                
                product = {
                    "Product Rank": f"#{rank}",
                    "Product Name": text,
                    "Brand Name": brand,
                    "Product Price": price,
                    "Discount Percentage": "",
                    "Original Price": "",
                    "Product Rating": rating,
                    "Number of Reviews": reviews,
                    "Product URL": product_url,
                    "Product Image URL": "",
                    "Availability": "Available",
                    "Category": "Sports & Fitness",
                    "ASIN": asin,
                    "Timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                }
                
                all_products.append(product)
                print(f"{rank}. {text[:60]}... {price if price else 'N/A'}")
                rank += 1
                
            except Exception as e:
                logging.error(f"Error: {e}")
                continue
        
        browser.close()
    
    return all_products

# =========================================================
# EXPORT
# =========================================================

def export_data(products):
    """Export to CSV and Excel"""
    
    if not products:
        print("No products found!")
        return
    
    print(f"\nExporting {len(products)} products...")
    
    df = pd.DataFrame(products)
    
    # Remove duplicates
    df.drop_duplicates(subset=['Product Name'], inplace=True)
    
    # CSV
    df.to_csv(OUTPUT_CSV, index=False, encoding='utf-8-sig')
    print(f"✓ CSV: {OUTPUT_CSV}")
    
    # Excel with formatting
    df.to_excel(OUTPUT_XLSX, index=False, engine='openpyxl')
    
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, PatternFill, Font
        
        wb = load_workbook(OUTPUT_XLSX)
        ws = wb.active
        
        # Header
        for cell in ws[1]:
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        
        # Column width
        for column in ws.columns:
            max_len = 0
            col = column[0].column_letter
            for cell in column:
                try:
                    max_len = max(max_len, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col].width = min(max_len + 2, 50)
        
        wb.save(OUTPUT_XLSX)
    except:
        pass
    
    print(f"✓ Excel: {OUTPUT_XLSX}")
    print(f"\n✓ Complete! {len(df)} products exported.")

# =========================================================
# MAIN
# =========================================================

if __name__ == "__main__":
    print("\n" + "="*60)
    print("Amazon Bestseller Scraper v3")
    print("="*60 + "\n")
    
    products = scrape_bestsellers()
    export_data(products)
