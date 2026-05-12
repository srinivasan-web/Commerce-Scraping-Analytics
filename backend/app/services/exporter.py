from __future__ import annotations

import json
import urllib.request
import zipfile
from io import BytesIO, StringIO

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.app.schemas import Product


ESSENTIAL_COLUMNS = [
    "Product Rank",
    "Rank Number",
    "Product Name",
    "Brand Name",
    "Product Price",
    "Discount Amount",
    "Coupon Amount",
    "Final Effective Price",
    "Product Rating Value",
    "Number of Reviews",
    "Overall Bought Count",
    "Bought Count Text",
    "Monthly Units Used",
    "Sales Velocity",
    "Demand Score",
    "Bought x Price Revenue",
    "Product URL",
    "Availability",
    "Category",
    "ASIN",
    "Timestamp",
]

VARIANT_COLUMNS = [
    "Parent ASIN",
    "Variant ASIN",
    "Variant Name",
    "Variant Type",
    "Variant Color",
    "Variant Size",
    "Variant Weight",
    "Variant Model",
    "Variant SKU",
    "Variant ASIN",
    "Variant Price",
    "Variant Original Price",
    "Variant Discount %",
    "Variant Discount Amount",
    "Variant Coupon Amount",
    "Variant Final Effective Price",
    "Variant Availability",
    "Variant Delivery",
    "Variant Warranty",
    "Variant Offer Text",
    "Bank Offers",
    "EMI Offers",
    "Cashback Offers",
    "Coupon Offers",
    "Partner Offers",
    "Exchange Offers",
    "Variant Bought Count Text",
    "Variant Overall Bought Count",
    "Variant Monthly Units Sold",
    "Variant Visible Bought Count",
    "Variant Units Source",
    "Variant Units Estimated",
    "Variant Estimated Revenue",
    "Revenue Score",
    "Demand Score",
    "Trend Score",
    "Popularity Score",
    "Bestseller Strength",
    "Rating Quality Score",
    "Sales Potential Score",
    "Price Competitiveness",
    "Discount Strength",
    "Offer Strength",
    "Price Category",
    "Discount Analysis",
]


def price_category(price: float) -> str:
    if price < 500:
        return "Budget"
    if price <= 2000:
        return "Mid Range"
    return "Premium"


def discount_analysis(discount: float) -> str:
    if discount >= 35:
        return "High Discount"
    if discount >= 15:
        return "Medium Discount"
    if discount > 0:
        return "Low Discount"
    return "No Discount"


def score(product: Product, key: str) -> float:
    return round(product.scores.get(key, 0), 2)


def product_row(product: Product) -> dict[str, object]:
    return {
        "Product Rank": product.product_rank,
        "Rank Number": product.rank_number,
        "Product Name": product.name,
        "Brand Name": product.brand_name,
        "Product Price": product.price_text,
        "Discount Amount": product.discount_amount,
        "Coupon Amount": product.coupon_amount,
        "Final Effective Price": product.final_effective_price or product.price,
        "Product Rating Value": product.rating,
        "Number of Reviews": product.reviews_text,
        "Overall Bought Count": product.visible_bought_count,
        "Bought Count Text": product.bought_count_text,
        "Monthly Units Used": product.units_sold,
        "Sales Velocity": product.sales_velocity,
        "Demand Score": product.demand_score or score(product, "demand_score"),
        "Bought x Price Revenue": product.revenue,
        "Product URL": product.product_url,
        "Availability": product.availability,
        "Category": product.category,
        "ASIN": product.parent_asin,
        "Timestamp": product.scraped_at.strftime("%Y-%m-%d %H:%M:%S"),
    }


def variant_row(product: Product) -> dict[str, object]:
    return {
        "Parent ASIN": product.parent_asin,
        "Variant ASIN": product.variant_asin,
        "Variant Name": product.variant_name or product.variant,
        "Variant Type": product.variant_type or "Detected",
        "Variant Color": product.color,
        "Variant Size": product.size,
        "Variant Weight": product.weight,
        "Variant Model": product.model,
        "Variant SKU": product.sku,
        "Variant ASIN": product.variant_asin,
        "Variant Price": product.price,
        "Variant Original Price": product.original_price,
        "Variant Discount %": product.discount,
        "Variant Discount Amount": product.discount_amount,
        "Variant Coupon Amount": product.coupon_amount,
        "Variant Final Effective Price": product.final_effective_price or product.price,
        "Variant Availability": product.availability,
        "Variant Delivery": product.delivery,
        "Variant Warranty": product.warranty,
        "Variant Offer Text": product.offers,
        "Bank Offers": product.bank_offers,
        "EMI Offers": product.emi_offers,
        "Cashback Offers": product.cashback_offers,
        "Coupon Offers": product.coupon_offers,
        "Partner Offers": product.partner_offers,
        "Exchange Offers": product.exchange_offers,
        "Variant Bought Count Text": product.bought_count_text,
        "Variant Overall Bought Count": product.visible_bought_count,
        "Variant Monthly Units Sold": product.units_sold,
        "Variant Visible Bought Count": product.visible_bought_count,
        "Variant Units Source": product.bought_count_source,
        "Variant Units Estimated": product.units_sold_estimated,
        "Variant Estimated Revenue": product.revenue,
        "Revenue Score": score(product, "revenue_score"),
        "Demand Score": score(product, "demand_score"),
        "Trend Score": score(product, "trend_score"),
        "Popularity Score": round((score(product, "demand_score") + score(product, "bestseller_score")) / 2, 2),
        "Bestseller Strength": score(product, "bestseller_score"),
        "Rating Quality Score": score(product, "rating_quality_score"),
        "Sales Potential Score": score(product, "performance_score"),
        "Price Competitiveness": score(product, "price_competitiveness"),
        "Discount Strength": product.discount,
        "Offer Strength": score(product, "offer_strength"),
        "Price Category": price_category(product.price),
        "Discount Analysis": discount_analysis(product.discount),
    }


def products_frame(products: list[Product]) -> pd.DataFrame:
    rows = [product_row(product) for product in products]
    return pd.DataFrame(rows, columns=ESSENTIAL_COLUMNS)


def variants_frame(products: list[Product]) -> pd.DataFrame:
    rows = [variant_row(product) for product in products]
    return pd.DataFrame(rows, columns=VARIANT_COLUMNS)


def export_csv(products: list[Product]) -> bytes:
    return products_frame(products).to_csv(index=False).encode("utf-8-sig")


def export_json(products: list[Product]) -> bytes:
    payload = {
        "parent_products": [product_row(product) for product in products],
        "product_variants": [variant_row(product) for product in products],
    }
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")


def export_txt(products: list[Product]) -> bytes:
    buffer = StringIO()
    for product in products:
        buffer.write(
            f"{product.product_rank} | {product.name} | {product.price_text} | "
            f"{product.visible_bought_count} scraped bought | {product.units_sold} monthly units "
            f"({product.bought_count_source or 'unknown'}) | revenue {product.revenue:.2f}\n"
        )
    return buffer.getvalue().encode("utf-8")


def write_frame(writer: pd.ExcelWriter, frame: pd.DataFrame, sheet_name: str) -> None:
    safe = frame.copy()
    if safe.empty:
        safe = pd.DataFrame({"Message": ["No rows available for this sheet"]})
    safe.to_excel(writer, index=False, sheet_name=sheet_name[:31])


def export_excel(products: list[Product]) -> bytes:
    parents = products_frame(products)
    variants = variants_frame(products)
    revenue = variants.sort_values("Variant Estimated Revenue", ascending=False)
    demand = variants.sort_values("Demand Score", ascending=False)
    offers = variants[["Parent ASIN", "Variant Name", "Variant Offer Text", "Bank Offers", "EMI Offers", "Cashback Offers", "Coupon Offers", "Partner Offers", "Exchange Offers", "Variant Discount %", "Offer Strength"]]
    delivery = pd.DataFrame(
        [
            {
                "Parent ASIN": product.parent_asin,
                "Variant Name": product.variant_name or product.variant,
                "Delivery": product.delivery,
                "Free Delivery": product.free_delivery,
                "Fast Delivery": product.fast_delivery,
                "Delivery Days": product.delivery_days,
                "Installation Available": product.installation_available,
            }
            for product in products
        ]
    )
    reviews = pd.DataFrame(
        [
            {
                "Product Name": product.name,
                "Variant Name": product.variant_name or product.variant,
                "Rating": product.rating,
                "Rating Count": product.reviews,
                "Review Count": product.reviews,
                "Positive Review %": product.positive_review_percent,
                "Negative Review %": product.negative_review_percent,
                "Review Keywords": ", ".join(product.review_keywords),
                "Review Quality Score": score(product, "rating_quality_score"),
            }
            for product in products
        ]
    )
    insights = pd.DataFrame(
        [
            {
                "Insight": "Revenue leader" if index == 0 else "Growth candidate",
                "Product Name": product.name,
                "Variant": product.variant_name or product.variant,
                "Signal Score": score(product, "performance_score"),
                "Recommendation": "Prioritize inventory, price monitoring, and offer tracking.",
            }
            for index, product in enumerate(products[:50])
        ]
    )
    forecast = pd.DataFrame(
        [
            {
                "Parent ASIN": product.parent_asin,
                "Variant": product.variant_name or product.variant,
                "Current Monthly Revenue": product.revenue,
                "Forecast Month 1": round(product.revenue * 1.06, 2),
                "Forecast Month 3": round(product.revenue * 1.18, 2),
                "Forecast Month 6": round(product.revenue * 1.42, 2),
            }
            for product in products
        ]
    )
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        write_frame(writer, parents, "Parent Products")
        write_frame(writer, variants, "Product Variants")
        write_frame(writer, revenue, "Revenue Analytics")
        write_frame(writer, revenue[["Parent ASIN", "Variant Name", "Variant Estimated Revenue", "Variant Price", "Variant Monthly Units Sold", "Revenue Score"]], "Variant Revenue")
        write_frame(writer, demand[["Parent ASIN", "Variant Name", "Demand Score", "Trend Score", "Popularity Score", "Variant Monthly Units Sold", "Sales Potential Score"]], "Demand Analytics")
        write_frame(writer, offers, "Offer Analytics")
        write_frame(writer, delivery, "Delivery Analytics")
        write_frame(writer, reviews, "Review Analytics")
        write_frame(writer, insights, "AI Insights")
        write_frame(writer, forecast, "Sales Forecasting")
        workbook = writer.book
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            headers = [str(cell.value or "") for cell in worksheet[1]]
            for column_index, column_cells in enumerate(worksheet.columns, start=1):
                width = min(max(len(str(cell.value or "")) for cell in column_cells) + 3, 60)
                worksheet.column_dimensions[get_column_letter(column_index)].width = max(width, 12)
                header = headers[column_index - 1]
                if header in {
                    "Product Price Value",
                    "Original Price Value",
                    "Estimated Monthly Revenue",
                    "Bought x Price Revenue",
                    "Product Total Revenue",
                    "Variant Price",
                    "Variant Original Price",
                    "Variant Estimated Revenue",
                }:
                    for cell in column_cells[1:]:
                        cell.number_format = '"INR" #,##0.00'
                elif header in {"Discount Percentage", "Variant Discount %"}:
                    for cell in column_cells[1:]:
                        cell.number_format = '0.00'
            for metric in ["Bought x Price Revenue", "Estimated Monthly Revenue", "Product Total Revenue", "Variant Estimated Revenue", "Demand Score", "Sales Potential Score"]:
                if metric in headers and worksheet.max_row > 2:
                    col = get_column_letter(headers.index(metric) + 1)
                    worksheet.conditional_formatting.add(
                        f"{col}2:{col}{worksheet.max_row}",
                        ColorScaleRule(
                            start_type="min",
                            start_color="FCA5A5",
                            mid_type="percentile",
                            mid_value=50,
                            mid_color="FDE68A",
                            end_type="max",
                            end_color="86EFAC",
                        ),
                    )
    output.seek(0)
    return output.read()


def export_images(products: list[Product]) -> bytes:
    output = BytesIO()
    image_urls = [product.image_url for product in products if product.image_url]
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("image_urls.txt", "\n".join(image_urls) or "No image URLs were captured for this export.")
        for index, product in enumerate([item for item in products if item.image_url][:50], start=1):
            try:
                with urllib.request.urlopen(product.image_url, timeout=8) as response:
                    content_type = response.headers.get("content-type", "")
                    extension = ".jpg"
                    if "png" in content_type:
                        extension = ".png"
                    elif "webp" in content_type:
                        extension = ".webp"
                    archive.writestr(f"images/{index:03d}-{product.id}{extension}", response.read())
            except Exception:
                continue
    output.seek(0)
    return output.read()
